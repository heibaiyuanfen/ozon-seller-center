"""Local, deterministic selection and pricing review. No network or AI calls."""
from decimal import Decimal, InvalidOperation, ROUND_FLOOR
from datetime import datetime
from pathlib import Path
from core import calculate_price_breakdown, calculate_shipping, extract_reference_article

DEFAULT_RULES = {
    "category_keyword": "", "min_sales": "100", "sales_period_days": "30",
    "min_price": "", "max_price": "", "max_weight": "",
    "min_roi": "60", "min_margin": "25", "label_fee": "2",
    "advertising_percent": "15", "cargo_loss_percent": "10", "max_age_days": "7",
}


def number(value, label, *, minimum=Decimal(0), maximum=None):
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ValueError(label + "缺失或不是有效数字") from None
    if not result.is_finite() or result < minimum or (maximum is not None and result >= maximum):
        raise ValueError(label + "超出允许范围")
    return result


def validate_rules(rules):
    result = {**DEFAULT_RULES, **rules}
    for key in ("min_sales", "min_price", "max_price", "max_weight"):
        if result[key] != "":
            number(result[key], key)
    for key in ("min_roi", "label_fee", "max_age_days"):
        number(result[key], key)
    for key in ("min_margin", "advertising_percent", "cargo_loss_percent"):
        number(result[key], key, maximum=Decimal(100))
    days = number(result["sales_period_days"], "销量周期", minimum=Decimal(1))
    if days != days.to_integral_value():
        raise ValueError("销量周期须为整数天")
    if result["min_price"] and result["max_price"] and Decimal(result["min_price"]) > Decimal(result["max_price"]):
        raise ValueError("最低市场价不能高于最高市场价")
    return result


def source_identity(item):
    source = str(item.get("listing_url") or item.get("source_url") or item.get("sku") or "").strip()
    return extract_reference_article(source) or str(item.get("sku") or source).strip()


def read_listing_history(path):
    """Use the local ledger even after completed queue rows have been cleared."""
    sources, offers = set(), set()
    if path is None or not Path(path).is_file():
        return sources, offers
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["产品台账"]
        rows = sheet.iter_rows(values_only=True)
        headers = {str(value): index for index, value in enumerate(next(rows))}
        for key in ("平台", "店铺配置ID", "货号", "Ozon来源链接"):
            if key not in headers:
                raise ValueError("产品台账缺少查重列：" + key)
        for row in rows:
            def cell(key):
                index = headers[key]
                return str(row[index] or "").strip() if index < len(row) else ""
            if cell("平台").casefold() != "ozon":
                continue
            shop = cell("店铺配置ID")
            if not shop:
                raise ValueError("Ozon 台账记录缺店铺配置ID，无法可靠查重，请先补齐")
            if cell("Ozon来源链接"):
                sources.add((source_identity({"source_url": cell("Ozon来源链接")}), shop))
            if cell("货号"):
                offers.add((cell("货号"), shop))
    finally:
        workbook.close()
    return sources, offers


def pricing_values(item, rules, fbp=False):
    result = {}
    for key in ("label_fee", "advertising_percent", "cargo_loss_percent"):
        result[key] = str(item.get(key) if str(item.get(key, "")).strip() else rules[key])
    result["label_fee"] = "0" if fbp else result["label_fee"]
    result["sales_commission_percent"] = str(item.get("sales_commission_percent") or "")
    result["target_roi"] = str(max(number(rules["min_roi"], "最低 ROI"), number(item.get("target_roi") or rules["min_roi"], "目标 ROI")))
    return result


def purchase_ceiling(price, weight_g, values, min_margin, fbp=False):
    price = number(price, "预计成交价", minimum=Decimal("0.01"))
    weight = number(weight_g, "包装毛重", minimum=Decimal("0.01")) / 1000
    commission = number(values["sales_commission_percent"], "标准佣金", maximum=Decimal(100)) / 100
    commission -= Decimal("0.01") if fbp else 0
    if commission <= 0:
        raise ValueError("有效佣金必须大于 0")
    advertising = number(values["advertising_percent"], "广告费率", maximum=Decimal(100)) / 100
    loss = number(values["cargo_loss_percent"], "货损率", maximum=Decimal(100)) / 100
    retained = 1 - commission - Decimal("0.02") - advertising
    if retained <= 0:
        raise ValueError("费用比例合计须小于 100%")
    net = price * retained - (1 + loss) * calculate_shipping(price, weight, values.get("shipping_exchange_rate", "12.5"))
    roi = number(values["target_roi"], "最低 ROI") / 100
    margin = number(min_margin, "最低销售利润率", maximum=Decimal(100)) / 100
    invested = min(net / (1 + loss + roi), (net - margin * price) / (1 + loss))
    ceiling = invested - number(values["label_fee"], "贴单费")
    return ceiling.quantize(Decimal("0.01"), rounding=ROUND_FLOOR)


def evaluate_product(item, rules, *, fbp=False, now=None, exchange_rate="12.5"):
    rules = validate_rules(rules)
    rejected, pending = [], []
    result = {"status": "待补充", "reason": "", "ceiling": "", "profit": "", "roi": "", "margin": "", "price": ""}
    keyword = str(rules["category_keyword"]).strip().casefold()
    if keyword and keyword not in str(item.get("category") or "").casefold():
        (rejected if item.get("category") else pending).append("类目不匹配" if item.get("category") else "缺类目")
    if rules["min_sales"] != "":
        try:
            sales = number(item.get("sales"), "销量")
            period = number(item.get("sales_period_days"), "销量统计周期", minimum=Decimal(1))
            if period != Decimal(rules["sales_period_days"]):
                pending.append("销量统计周期与规则不一致")
            elif sales < Decimal(rules["min_sales"]):
                rejected.append("销量未达标")
        except ValueError as error:
            pending.append(str(error))
    try:
        market = number(item.get("market_price"), "市场参考价", minimum=Decimal("0.01"))
        for key, fail in (("min_price", market < Decimal(rules["min_price"] or 0)), ("max_price", market > Decimal(rules["max_price"] or market))):
            if rules[key] and fail:
                rejected.append("市场价超出筛选区间")
        # A user-reviewed intended selling price wins; never substitute an ROI-derived higher price.
        price = number(item.get("review_price") or market, "预计成交价", minimum=Decimal("0.01"))
        result["price"] = str(price)
        weight = number(item.get("weight"), "包装毛重", minimum=Decimal("0.01"))
        if rules["max_weight"] and weight > Decimal(rules["max_weight"]):
            rejected.append("包装毛重超限")
        values = pricing_values(item, rules, fbp)
        values["shipping_exchange_rate"] = str(exchange_rate)
        ceiling = purchase_ceiling(price, weight, values, rules["min_margin"], fbp)
        result["ceiling"] = str(ceiling)
        if ceiling <= 0:
            rejected.append("该成交价没有采购预算")
        purchase = number(item.get("purchase_cost"), "采购到手成本", minimum=Decimal("0.01"))
        breakdown = calculate_price_breakdown(price, purchase, values["label_fee"], weight / 1000,
            values["sales_commission_percent"], exchange_rate=exchange_rate, sales_commission_discount_percent="1" if fbp else "0",
            advertising_percent=values["advertising_percent"], cargo_loss_percent=values["cargo_loss_percent"])
        margin = breakdown.profit / price * 100
        result.update(profit=f"{breakdown.profit:.2f}", roi=f"{breakdown.roi_percent:.2f}", margin=f"{margin:.2f}")
        if breakdown.roi_percent < Decimal(values["target_roi"]) or margin < Decimal(rules["min_margin"]):
            rejected.append("预计成交价下利润未达标")
    except ValueError as error:
        pending.append(str(error))
    if not str(item.get("supplier_1688_url") or "").startswith(("https://", "http://")):
        pending.append("缺有效采购链接")
    if not str(item.get("supplier_sku") or "").strip():
        pending.append("缺采购规格/SKU")
    current = now or datetime.now()
    for flag, label in (("supplier_confirmed", "同款规格/交期确认"), ("economics_confirmed", "成交价/到手成本/包装参数确认")):
        if str(item.get(flag) or "") != "1":
            pending.append("待" + label)
        else:
            try:
                age = (current - datetime.fromisoformat(str(item.get("review_confirmed_at") or ""))).total_seconds() / 86400
                if age < 0 or age > Decimal(rules["max_age_days"]):
                    pending.append("确认数据已过期，请重新审核")
            except ValueError:
                pending.append("缺确认时间，请重新审核")
    for field, label in (("net_weight", "净重"), ("depth", "包装长"), ("width", "包装宽"), ("height", "包装高")):
        try:
            number(item.get(field), label, minimum=Decimal("0.01"))
        except ValueError as error:
            pending.append(str(error))
    if not str(item.get("model_name") or "").strip():
        pending.append("缺人工型号名称")
    if not source_identity(item):
        pending.append("缺来源 SKU/链接")
    result["status"] = "淘汰" if rejected else "待补充" if pending else "通过"
    result["reason"] = "；".join(dict.fromkeys(rejected + pending)) or "规则与人工货源确认通过，可预检入队"
    return result
