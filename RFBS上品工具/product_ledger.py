from __future__ import annotations

import os
import time
import threading
import uuid
import json
import shutil
from collections import defaultdict
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SHEET_TITLE = "产品台账"
TABLE_NAME = "ProductLedgerTable"


@dataclass(frozen=True)
class LedgerColumn:
    key: str
    title: str
    width: float
    kind: str = "text"


LEDGER_COLUMNS = (
    LedgerColumn("description", "商品描述", 50),
    LedgerColumn("tags", "搜索标签", 45),
    LedgerColumn("attributes_json", "普通属性JSON", 50),
    LedgerColumn("complex_attributes_json", "组合属性JSON", 50),
    LedgerColumn("shipping_exchange_rate", "上架核价汇率", 16),
    LedgerColumn("snapshot_version", "上架快照版本", 14),
    LedgerColumn("updated_at", "更新时间", 20, "datetime"),
    LedgerColumn("platform", "平台", 14),
    LedgerColumn("shop_name", "上品店铺", 18),
    LedgerColumn("listing_mode", "上品模式", 12),
    LedgerColumn("offer_id", "货号", 22),
    LedgerColumn("supplier_1688_url", "1688采购链接", 38, "link"),
    LedgerColumn("ozon_source_url", "Ozon来源链接", 38, "link"),
    LedgerColumn("ozon_product_id", "Ozon商品ID", 18),
    LedgerColumn("wb_source_url", "WB来源链接", 38, "link"),
    LedgerColumn("wb_source_nm_id", "WB来源商品ID", 18),
    LedgerColumn("wb_nm_id", "WB商品ID(nmID)", 18),
    LedgerColumn("wb_chrt_id", "WB规格ID(chrtID)", 18),
    LedgerColumn("wb_subject_id", "WB子类目ID", 16),
    LedgerColumn("title", "商品标题", 36),
    LedgerColumn("model_name", "型号名称", 24),
    LedgerColumn("price", "售价", 13, "money"),
    LedgerColumn("old_price", "折扣前价格", 13, "money"),
    LedgerColumn("currency_code", "币种", 9),
    LedgerColumn("pricing_mode", "核价模式", 11),
    LedgerColumn("purchase_cost", "采购成本", 13, "money"),
    LedgerColumn("label_fee", "贴单费", 11, "money"),
    LedgerColumn("target_roi_percent", "目标ROI", 11, "percent"),
    LedgerColumn("actual_roi_percent", "实际ROI", 11, "percent"),
    LedgerColumn("target_profit_margin_percent", "目标利润率", 12, "percent"),
    LedgerColumn("actual_profit_margin_percent", "实际利润率", 12, "percent"),
    LedgerColumn("discount_percent", "WB折扣", 10, "percent"),
    LedgerColumn("advertising_percent", "广告费率", 11, "percent"),
    LedgerColumn("cargo_loss_percent", "货损率", 11, "percent"),
    LedgerColumn("raw_sales_commission_percent", "填写标准佣金", 15, "percent"),
    LedgerColumn("sales_commission_percent", "核价佣金", 13, "percent"),
    LedgerColumn("shipping", "运费", 12, "money"),
    LedgerColumn("cny_rub_rate", "CNY兑RUB汇率", 14, "number"),
    LedgerColumn("cross_border_shipping_cny", "WB跨境运费(CNY)", 18, "money"),
    LedgerColumn("first_mile_shipping_cny", "WB头程运费(CNY)", 18, "money"),
    LedgerColumn("warehouse_operation_fee_cny", "海外仓操作费(CNY)", 19, "money"),
    LedgerColumn("tail_delivery_rub", "WB尾程费(RUB)", 16, "money"),
    LedgerColumn("volume_liters", "包装体积(L)", 13, "number"),
    LedgerColumn("sales_commission", "销售佣金", 12, "money"),
    LedgerColumn("logistics_commission", "物流佣金", 12, "money"),
    LedgerColumn("advertising", "广告费", 12, "money"),
    LedgerColumn("cargo_loss", "货损", 12, "money"),
    LedgerColumn("profit", "利润", 12, "money"),
    LedgerColumn("net_weight_g", "商品净重(g)", 14, "number"),
    LedgerColumn("gross_weight_g", "包装毛重(g)", 14, "number"),
    LedgerColumn("depth_mm", "包装长度(mm)", 15, "number"),
    LedgerColumn("width_mm", "包装宽度(mm)", 15, "number"),
    LedgerColumn("height_mm", "包装高度(mm)", 15, "number"),
    LedgerColumn("category_display", "Ozon类目", 34),
    LedgerColumn("category_id", "类目ID", 16),
    LedgerColumn("type_id", "类型ID", 16),
    LedgerColumn("warehouse_id", "RFBS仓库ID", 20),
    LedgerColumn("stock", "可售库存", 11, "integer"),
    LedgerColumn("import_task_id", "Ozon导入任务ID", 18),
    LedgerColumn("status", "状态", 22),
    LedgerColumn("shop_config_id", "店铺配置ID", 18),
    LedgerColumn("product_group_id", "任务产品组ID", 28),
    LedgerColumn("base_offer_id", "原始货号", 24),
    LedgerColumn("backup_product_id", "备份产品ID", 34),
    LedgerColumn("backup_variant_id", "备份规格ID", 34),
    LedgerColumn("dimension_unit", "台账尺寸单位", 14),
    LedgerColumn("weight_unit", "台账重量单位", 14),
    LedgerColumn("primary_image_url", "主图链接", 50, "link"),
    LedgerColumn("additional_image_urls", "附图链接", 60),
    LedgerColumn("image_urls_json", "全部图片链接JSON", 65),
)


_LEDGER_LOCK = threading.RLock()


def image_link_fields(primary=None, images=()):
    values = [primary, *images]
    urls = list(dict.fromkeys(
        value.strip() for value in values
        if isinstance(value, str) and value.strip().startswith(('https://', 'http://'))
    ))
    return dict(primary_image_url=urls[0] if urls else '',
                additional_image_urls='\n'.join(urls[1:]),
                image_urls_json=json.dumps(urls, ensure_ascii=False))


def ledger_image_urls(row):
    raw = row.get('image_urls_json') or '[]'
    values = json.loads(raw) if isinstance(raw, str) else raw
    if not isinstance(values, list) or any(not isinstance(value, str) for value in values):
        raise ValueError('全部图片链接JSON必须是字符串数组')
    additional = str(row.get('additional_image_urls') or '').splitlines()
    return list(dict.fromkeys(value.strip() for value in
                [row.get('primary_image_url') or '', *values, *additional] if value.strip()))


def backfill_ledger_image_links(path, submission_dir):
    """Fill missing image cells only from an unambiguous successful submission."""
    path, submission_dir = Path(path).resolve(), Path(submission_dir)
    updated = 0
    backup = ''
    with _LEDGER_LOCK, _process_file_lock(path):
        workbook, sheet, _ = _load_or_create_workbook(path)
        try:
            headers = _ensure_headers(sheet)
            rows = [(index, {column.key: sheet.cell(index, headers[column.title]).value
                             for column in LEDGER_COLUMNS}) for index in range(2, sheet.max_row + 1)]
            owners = defaultdict(set)
            for _, row in rows:
                if str(row.get('platform') or '').lower() == 'ozon':
                    owners[(str(row.get('offer_id') or ''), str(row.get('import_task_id') or ''))].add(
                        str(row.get('shop_config_id') or ''))
            first_updated = None
            for index, row in rows:
                if str(row.get('platform') or '').lower() != 'ozon' or ledger_image_urls(row):
                    continue
                offer, task, shop = (str(row.get(key) or '') for key in
                                     ('offer_id', 'import_task_id', 'shop_config_id'))
                if not offer or not task or not shop or '/' in offer or '\\' in offer:
                    continue
                source = submission_dir / ('submitted-' + offer + '.json')
                try:
                    document = json.loads(source.read_text(encoding='utf-8-sig'))
                    item = document.get('item') or {}
                    if not document.get('submitted') or item.get('offer_id') != offer or str(document.get('task_id') or '') != task:
                        continue
                    source_shop = str(document.get('shop_config_id') or '')
                    if (source_shop and source_shop != shop) or (not source_shop and len(owners[(offer, task)]) != 1):
                        continue
                    entries = (document.get('import_result') or {}).get('result', {}).get('items', [])
                    if not any(entry.get('offer_id') == offer and entry.get('status') == 'imported'
                               and not entry.get('errors') for entry in entries):
                        continue
                    fields = image_link_fields(item.get('primary_image'), item.get('images') or [])
                    if not fields['primary_image_url']:
                        continue
                except (OSError, ValueError, TypeError, AttributeError):
                    continue
                for key, value in fields.items():
                    column = next(column for column in LEDGER_COLUMNS if column.key == key)
                    sheet.cell(index, headers[column.title], value)
                updated += 1
                first_updated = first_updated or index
            if updated:
                directory = path.parent / '台账备份'
                directory.mkdir(parents=True, exist_ok=True)
                destination = directory / f'{path.stem}-补图前-{uuid.uuid4().hex[:12]}.xlsx'
                shutil.copy2(path, destination)
                backup = str(destination)
                _style_worksheet(sheet, headers, first_updated)
                _ensure_table(sheet, headers)
                _save_atomic(workbook, path)
        finally:
            workbook.close()
    return dict(updated=updated, backup=backup)


def listing_snapshot(item, content):
    """Only submission data; never credentials, image files, URLs or local image paths."""
    excluded = {"images", "primary_image", "images360", "color_image", "image", "pictures"}
    def clean(value):
        if isinstance(value, dict):
            return {key: clean(part) for key, part in value.items() if key not in excluded}
        if isinstance(value, list):
            return [clean(part) for part in value]
        return value
    return json.dumps({"version": 1, "item": clean(item), "content": content}, ensure_ascii=False)


def write_snapshot(workbook, record):
    raw = record.get("listing_snapshot_json")
    if not raw:
        return
    name = "上架信息快照"
    sheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
    if sheet.cell(1, 1).value is None:
        sheet.append(["平台", "店铺配置ID", "货号", "分片序号", "JSON内容"])
    identity = tuple(str(record.get(key) or "") for key in ("platform", "shop_config_id", "offer_id"))
    for row in range(sheet.max_row, 1, -1):
        if tuple(str(sheet.cell(row, col).value or "") for col in (1, 2, 3)) == identity:
            sheet.delete_rows(row)
    # Stay below Excel's 32767-character cell limit without losing full attributes/content.
    for index, start in enumerate(range(0, len(raw), 12000), start=1):
        sheet.append([*identity, index, raw[start:start + 12000]])
        sheet.cell(sheet.max_row, 5).data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["E"].width = 90


def read_listing_snapshot(path, platform, shop_config_id, offer_id):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "上架信息快照" not in workbook.sheetnames:
            return None
        identity = tuple(map(str, (platform, shop_config_id, offer_id)))
        parts = [(int(row[3]), row[4]) for row in workbook["上架信息快照"].iter_rows(min_row=2, values_only=True)
                 if tuple(str(value or "") for value in row[:3]) == identity]
        return json.loads("".join(part for _, part in sorted(parts))) if parts else None
    finally:
        workbook.close()
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
_BODY_FONT = Font(name="Microsoft YaHei", size=10, color="1F2933")
_LINK_FONT = Font(name="Microsoft YaHei", size=10, color="0563C1", underline="single")
_THIN_GRAY = Side(style="thin", color="D9E2F3")
_BODY_BORDER = Border(bottom=_THIN_GRAY)


def _safe_text(value) -> str:
    text = str(value or "").strip()
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _decimal_value(value, title: str):
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"产品台账字段“{title}”必须是有效数字：{value}") from error
    if not number.is_finite():
        raise ValueError(f"产品台账字段“{title}”必须是有限数字：{value}")
    return float(number)


def _cell_value(column: LedgerColumn, value):
    if column.kind in {"text", "link"}:
        return _safe_text(value)
    if column.kind == "datetime":
        if isinstance(value, datetime):
            return value.replace(microsecond=0)
        if value is None or str(value).strip() == "":
            return datetime.now().replace(microsecond=0)
        text = str(value).strip()
        try:
            return datetime.fromisoformat(text)
        except ValueError:
            return _safe_text(text)
    if column.kind == "integer":
        number = _decimal_value(value, column.title)
        return int(number) if number is not None else None
    number = _decimal_value(value, column.title)
    if column.kind == "percent" and number is not None:
        return number / 100
    return number


def _load_or_create_workbook(path: Path):
    if not path.exists():
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = SHEET_TITLE
        workbook.properties.title = "Ozon RFBS 产品台账"
        workbook.properties.subject = "自动上架完成记录"
        return workbook, worksheet, True
    try:
        workbook = load_workbook(path)
    except PermissionError as error:
        raise PermissionError(f"无法读取产品台账，请先关闭 Excel：{path}") from error
    except Exception as error:
        raise ValueError(f"无法读取现有产品台账，请检查文件是否损坏：{path}\n{error}") from error
    worksheet = workbook[SHEET_TITLE] if SHEET_TITLE in workbook.sheetnames else workbook.create_sheet(SHEET_TITLE)
    return workbook, worksheet, False


def _ensure_headers(worksheet) -> dict[str, int]:
    header_by_title = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value is not None and str(cell.value).strip()
    }
    legacy_commission_title = "RFBS销售佣金"
    if "核价佣金" not in header_by_title and legacy_commission_title in header_by_title:
        column_index = header_by_title.pop(legacy_commission_title)
        worksheet.cell(1, column_index, "核价佣金")
        header_by_title["核价佣金"] = column_index
    legacy_manual_commission_title = "Ozon原始RFBS佣金"
    if "填写标准佣金" not in header_by_title and legacy_manual_commission_title in header_by_title:
        column_index = header_by_title.pop(legacy_manual_commission_title)
        worksheet.cell(1, column_index, "填写标准佣金")
        header_by_title["填写标准佣金"] = column_index
    for column in LEDGER_COLUMNS:
        if column.title not in header_by_title:
            index = len(header_by_title) + 1
            while worksheet.cell(1, index).value not in (None, ""):
                index += 1
            worksheet.cell(1, index, column.title)
            header_by_title[column.title] = index
    return header_by_title


def _matching_row(worksheet, header_by_title: dict[str, int], record: dict) -> int | None:
    offer_id = str(record.get("offer_id") or "").strip()
    shop_config_id = str(record.get("shop_config_id") or "").strip()
    offer_column = header_by_title["货号"]
    shop_column = header_by_title["店铺配置ID"]
    for row in range(2, worksheet.max_row + 1):
        existing_offer = str(worksheet.cell(row, offer_column).value or "").strip().lstrip("'")
        if existing_offer != offer_id:
            continue
        existing_shop = str(worksheet.cell(row, shop_column).value or "").strip().lstrip("'")
        if shop_config_id and existing_shop and existing_shop != shop_config_id:
            continue
        return row
    return None


def _style_worksheet(worksheet, header_by_title: dict[str, int], target_row: int):
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "C2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(max(header_by_title.values()))}{worksheet.max_row}"
    worksheet.row_dimensions[1].height = 28
    worksheet.row_dimensions[target_row].height = 22
    worksheet.sheet_properties.tabColor = "1F4E78"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    for column in LEDGER_COLUMNS:
        column_index = header_by_title[column.title]
        letter = get_column_letter(column_index)
        worksheet.column_dimensions[letter].width = column.width
        header = worksheet.cell(1, column_index)
        header.fill = _HEADER_FILL
        header.font = _HEADER_FONT
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row, column_index)
            cell.font = _BODY_FONT
            cell.border = _BODY_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=column.kind == "text")
            if column.kind in {"money", "number"}:
                cell.number_format = "#,##0.00"
            elif column.kind == "percent":
                cell.number_format = "0.00%"
            elif column.kind == "integer":
                cell.number_format = "0"
            elif column.kind == "datetime":
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            else:
                cell.number_format = "@"
            if column.kind == "link" and str(cell.value or "").startswith(("http://", "https://")):
                cell.hyperlink = str(cell.value)
                cell.font = _LINK_FONT
            elif column.kind == "link":
                cell.hyperlink = None

    worksheet.column_dimensions[get_column_letter(header_by_title["店铺配置ID"])].hidden = True


def _ensure_table(worksheet, header_by_title: dict[str, int]):
    if worksheet.max_row < 2:
        return
    last_column = max(header_by_title.values())
    for index in range(1, last_column + 1):
        if worksheet.cell(1, index).value in (None, ""):
            worksheet.cell(1, index, f"字段{index}")
    reference = f"A1:{get_column_letter(last_column)}{worksheet.max_row}"
    if worksheet.tables.get(TABLE_NAME) is not None:
        del worksheet.tables[TABLE_NAME]
    table = Table(displayName=TABLE_NAME, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _save_atomic(workbook, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.stem}-{uuid.uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    except PermissionError as error:
        raise PermissionError(
            f"无法更新产品台账，请关闭正在打开的 Excel 文件后重试：{path}"
        ) from error
    finally:
        try:
            temporary_path.unlink(missing_ok=True)
        except OSError:
            pass


@contextmanager
def _process_file_lock(path: Path, timeout: float = 60):
    """Serialize workbook updates across parallel listing processes on Windows."""
    lock_path = path.with_suffix(path.suffix + ".lock")
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    handle = open(lock_path, "a+b")
    deadline = time.monotonic() + timeout
    try:
        if os.name == "nt":
            import msvcrt
            while True:
                try:
                    handle.seek(0)
                    if handle.tell() == 0 and lock_path.stat().st_size == 0:
                        handle.write(b"0")
                        handle.flush()
                    handle.seek(0)
                    msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError:
                    if time.monotonic() >= deadline:
                        raise TimeoutError("等待 Excel 产品台账写入锁超时")
                    time.sleep(0.1)
        yield
    finally:
        if os.name == "nt":
            try:
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            except OSError:
                pass
        handle.close()
        try:
            lock_path.unlink(missing_ok=True)
        except OSError:
            pass


def upsert_product_ledger(path: str | Path, record: dict) -> dict:
    """Create or update one listing row, keyed by shop and Offer ID."""
    ledger_path = Path(path).expanduser().resolve()
    offer_id = str(record.get("offer_id") or "").strip()
    if not offer_id:
        raise ValueError("写入产品台账时货号不能为空")

    with _LEDGER_LOCK, _process_file_lock(ledger_path):
        workbook, worksheet, workbook_created = _load_or_create_workbook(ledger_path)
        header_by_title = _ensure_headers(worksheet)
        target_row = _matching_row(worksheet, header_by_title, record)
        action = "updated"
        if target_row is None:
            target_row = max(2, worksheet.max_row + 1)
            action = "created" if workbook_created and target_row == 2 else "added"

        values = dict(record)
        values["updated_at"] = values.get("updated_at") or datetime.now()
        for column in LEDGER_COLUMNS:
            if column.key in {'primary_image_url', 'additional_image_urls', 'image_urls_json'} and column.key not in record:
                continue
            cell = worksheet.cell(target_row, header_by_title[column.title])
            cell.value = _cell_value(column, values.get(column.key))

        _style_worksheet(worksheet, header_by_title, target_row)
        _ensure_table(worksheet, header_by_title)
        write_snapshot(workbook, record)
        _save_atomic(workbook, ledger_path)

    return {
        "path": str(ledger_path),
        "sheet": SHEET_TITLE,
        "row": target_row,
        "action": action,
        "offer_id": offer_id,
    }
