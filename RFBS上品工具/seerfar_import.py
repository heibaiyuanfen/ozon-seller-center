from __future__ import annotations

import re
from pathlib import Path
from typing import Any


REQUIRED_HEADERS = {"Title", "Listing URL", "SKU"}

# Seerfar exports localized column names according to the UI language. Keep one
# canonical set internally so old English exports and current Chinese exports
# follow exactly the same parsing path.
HEADER_ALIASES = {
    "排名": "NO.",
    "主图": "Image",
    "标题": "Title",
    "详情页地址": "Listing URL",
    "品牌": "Brand",
    "类目": "Categories",
    "售价": "Price",
    "销量": "Sales",
    "销售额": "Revenue",
    "毛利率": "Gross Margin",
    "评分": "Ratings",
    "评论数": "NO. Ratings",
    "店铺": "Shop",
    "卖家类型": "Seller type",
    "配送方式": "Fulfillment",
    "重量": "Weight",
    "上架时间": "Launch Age",
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _canonical_header(value: Any) -> str:
    header = _text(value).replace("\u00a0", " ").strip()
    return HEADER_ALIASES.get(header, header)


def _number_text(value: Any) -> str:
    raw = _text(value).replace("\u00a0", " ")
    match = re.search(r"-?\d+(?:[.,]\d+)?", raw.replace(" ", ""))
    return match.group(0).replace(",", ".") if match else ""


def _image_url_from_formula(value: Any) -> str:
    raw = _text(value)
    match = re.search(r'IMAGE\(\s*"([^"]+)"', raw, flags=re.I)
    return match.group(1) if match else ""


def load_seerfar_workbook(path: str | Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("缺少 Excel 读取组件 openpyxl") from error

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError("Excel 文件不存在：" + str(source))
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except Exception as error:
        raise ValueError("无法读取 Excel 文件，请确认它是有效的 .xlsx 文档") from error
    try:
        sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as error:
            raise ValueError("Excel 没有数据") from error
        headers = [_canonical_header(value) for value in raw_headers]
        missing = sorted(REQUIRED_HEADERS - set(headers))
        if missing:
            raise ValueError("不是支持的 Seerfar 选品表，缺少列：" + "、".join(missing))
        index = {name: position for position, name in enumerate(headers) if name}

        def cell(row, name):
            position = index.get(name)
            return row[position] if position is not None and position < len(row) else None

        products = []
        seen = set()
        for excel_row, row in enumerate(rows, start=2):
            title = _text(cell(row, "Title"))
            listing_url = _text(cell(row, "Listing URL"))
            sku = _text(cell(row, "SKU"))
            if not (title or listing_url or sku):
                continue
            identity = sku or listing_url or f"row-{excel_row}"
            if identity.casefold() in seen:
                continue
            seen.add(identity.casefold())
            image_url = _text(cell(row, "Image"))
            # Seerfar has two columns named Image: the first is a formula and
            # the second is the direct URL. Prefer the direct URL by position.
            image_positions = [i for i, name in enumerate(headers) if name == "Image"]
            if len(image_positions) > 1 and image_positions[1] < len(row):
                image_url = _text(row[image_positions[1]])
            if not image_url and image_positions:
                image_url = _image_url_from_formula(row[image_positions[0]])
            products.append({
                "id": identity,
                "excel_row": str(excel_row),
                "no": _text(cell(row, "NO.")),
                "image_url": image_url,
                "title": title,
                "listing_url": listing_url,
                "sku": sku,
                "brand": _text(cell(row, "Brand")),
                "category": _text(cell(row, "Categories")),
                "market_price": _number_text(cell(row, "Price")),
                "sales": _number_text(cell(row, "Sales")),
                "revenue": _number_text(cell(row, "Revenue")),
                "gross_margin": _number_text(cell(row, "Gross Margin")),
                "ratings": _number_text(cell(row, "Ratings")),
                "rating_count": _number_text(cell(row, "NO. Ratings")),
                "shop": _text(cell(row, "Shop")),
                "seller_type": _text(cell(row, "Seller type")),
                "fulfillment": _text(cell(row, "Fulfillment")),
                "source_weight": _number_text(cell(row, "Weight")),
                "launch_age": _text(cell(row, "Launch Age")),
            })
        if not products:
            raise ValueError("Excel 中没有可导入的商品行")
        return products
    finally:
        workbook.close()


MAPPING_FIELDS = (
    "supplier_1688_url", "offer_id", "model_name", "purchase_cost", "price",
    "old_price", "target_roi", "sales_commission_percent", "net_weight", "weight",
    "depth", "width", "height", "stock", "notes",
)


def merge_product_mappings(products: list[dict], mappings: dict) -> list[dict]:
    merged = []
    for product in products:
        item = dict(product)
        saved = mappings.get(str(product.get("id") or ""), {}) if isinstance(mappings, dict) else {}
        for field in MAPPING_FIELDS:
            item[field] = _text(saved.get(field)) if isinstance(saved, dict) else ""
        if not item["price"]:
            item["price"] = item.get("market_price", "")
        if not item["weight"] and item.get("source_weight"):
            item["weight"] = item["source_weight"]
        if not item["net_weight"] and item.get("source_weight"):
            item["net_weight"] = item["source_weight"]
        merged.append(item)
    return merged


def mapping_payload(products: list[dict]) -> dict[str, dict[str, str]]:
    return {
        str(item.get("id") or ""): {
            field: _text(item.get(field)) for field in MAPPING_FIELDS
        }
        for item in products if str(item.get("id") or "").strip()
    }
