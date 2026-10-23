"""Legacy ordinary-draft importer, retained for migration diagnostics only."""
import hashlib
import json
import uuid
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from core import atomic_write_json, load_json
from erp_backup import base_url, no_credentials, now, outbox_lock
from product_ledger import LEDGER_COLUMNS, SHEET_TITLE


def read_ledger(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_TITLE not in workbook.sheetnames:
            raise ValueError('找不到“产品台账”工作表')
        rows = workbook[SHEET_TITLE].iter_rows(values_only=True)
        headers = next(rows, ())
        columns = {column.title: column.key for column in LEDGER_COLUMNS}
        keys = [columns.get(str(value).strip()) for value in headers]
        if 'offer_id' not in keys or 'title' not in keys:
            raise ValueError('产品台账必须包含“货号”和“商品标题”表头')
        for number, row in enumerate(rows, 2):
            record = {key: (value.isoformat() if isinstance(value, (date, datetime)) else value)
                      for key, value in zip(keys, row) if key}
            if any(value is not None and str(value).strip() for value in record.values()):
                yield number, record
    finally:
        workbook.close()


def draft_from_row(row, marker):
    def text(key):
        value = row.get(key)
        return '' if value is None else str(value).strip()
    def number(key, positive=False, integer=False):
        raw = row.get(key)
        if raw is None or str(raw).strip() == '':
            return None
        value = Decimal(str(raw))
        if not value.is_finite() or value < 0 or (positive and value == 0) or value > 1000000000:
            raise ValueError(f'{key} 不是有效非负数或超出范围')
        if integer and value != int(value):
            raise ValueError(f'{key} 必须是整数')
        return int(value) if integer else float(value)
    if not text('offer_id') or not text('title'):
        raise ValueError('普通草稿接口仍要求货号和商品标题，不能留空')
    payload = dict(sku=text('offer_id'), name=text('title'), source_type='follow', status='draft',
        group_code=text('model_name'), source_url=text('ozon_source_url') or text('supplier_1688_url'),
        supplier='1688' if text('supplier_1688_url') else '', currency='CNY',
        purchase_price_cny=number('purchase_cost'), domestic_delivery_cny=None,
        # Spreadsheet measurements have no verified provenance; retain them only in the snapshot.
        package_length_mm=None, package_width_mm=None, package_height_mm=None, package_weight_g=None,
        description_category_id=number('category_id', positive=True, integer=True),
        type_id=number('type_id', positive=True, integer=True), category_display=text('category_display'),
        notes='产品台账导入草稿；图片、属性和包装事实尚未核验。',
        source_snapshot={'ledger_import': {'id': marker, 'row': row, 'images_verified': False, 'copy_ready': False}},
        variants=[dict(sku=text('offer_id'), specification=text('model_name'),
                       purchase_price_cny=number('purchase_cost'), properties={'ledger_import_id': marker})])
    no_credentials(payload)
    return payload


def import_ledger(path, root, client, log=lambda text: None):
    root = Path(root) / 'erp-ledger'
    target = base_url(client.url)
    report = {'started_at': now(), 'file': str(Path(path).resolve()), 'items': []}
    client.login()
    with outbox_lock(root):
        state_path = root / 'state.json'
        state = load_json(state_path, {'records': {}})
        for row_number, row in read_ledger(path):
            result = {'row': row_number, 'offer_id': row.get('offer_id')}
            try:
                locator = [target, str(row.get('platform') or ''), str(row.get('shop_config_id') or row.get('shop_name') or ''), str(row.get('offer_id') or '')]
                key = hashlib.sha256(json.dumps(locator, ensure_ascii=False).encode()).hexdigest()
                record = state['records'].setdefault(key, {'id': uuid.uuid4().hex})
                payload = draft_from_row(row, record['id'])
                digest = hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
                if record.get('status') == 'completed':
                    if record.get('digest') != digest:
                        raise ValueError('该行已导入后发生变化，为保护产品库人工编辑，不自动覆盖')
                    result.update(status='duplicate', product_id=record['product_id'])
                elif record.get('status') == 'sending':
                    raise ValueError('上次创建结果不确定，禁止自动重复创建；请先在产品库核对并联系处理')
                else:
                    atomic_write_json(root / (record['id'] + '.request.json'), payload)
                    record.update(status='sending', digest=digest)
                    atomic_write_json(state_path, state)
                    # This older endpoint has no idempotency key. Never blindly retry POST.
                    response = client.session.request('POST', target + '/api/v1/catalog/products',
                        data=json.dumps(payload, ensure_ascii=False, allow_nan=False).encode('utf-8'),
                        headers={'Content-Type': 'application/json; charset=utf-8'},
                        timeout=(5, 60), allow_redirects=False)
                    if response.status_code != 201:
                        if 400 <= response.status_code < 500:
                            record['status'] = 'rejected'
                            atomic_write_json(state_path, state)
                        raise ValueError(f'草稿接口返回 HTTP {response.status_code}，请核对权限和字段')
                    receipt = response.json()
                    if not receipt.get('id'):
                        raise ValueError('草稿接口缺少产品 ID，结果待核对')
                    record.update(status='completed', product_id=receipt['id'])
                    atomic_write_json(state_path, state)
                    result.update(status='created', product_id=receipt['id'])
            except Exception as error:
                result.update(status='failed', error=str(error))
            report['items'].append(result)
            log(f"台账第 {row_number} 行：{result['status']}" + (f"；{result['error']}" if result.get('error') else ''))
            atomic_write_json(root / 'latest-report.json', report)
    totals = {status: sum(item['status'] == status for item in report['items']) for status in ('created', 'duplicate', 'failed')}
    log(f"台账上传完成：新增 {totals['created']}，已导入 {totals['duplicate']}，待处理 {totals['failed']}；报告：{root / 'latest-report.json'}")
    return report
