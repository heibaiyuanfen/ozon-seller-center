import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

from seerfar_import import load_seerfar_workbook, mapping_payload, merge_product_mappings


class SeerfarImportTests(unittest.TestCase):
    def _workbook(self, path: Path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append([
            "NO.", "Image", "Image", "Title", "Listing URL", "SKU", "Brand", "Categories",
            "Price", "Sales", "Revenue", "Gross Margin", "Ratings", "NO. Ratings", "Shop",
            "Seller type", "Fulfillment", "Weight", "Launch Age",
        ])
        sheet.append([
            1, '=IMAGE("https://ir.ozone.ru/a.jpg")', "https://ir.ozone.ru/a.jpg", "Lamp",
            "https://www.ozon.ru/product/1234567890", 1234567890, "Brand", "Desk Lamp",
            "2536₽", 2896, "7495854₽", "13.5%", 4.9, 1840, "Shop", "Domestic seller",
            "FBO", "950.0 g", "1 Month",
        ])
        workbook.save(path)

    def test_reads_seerfar_rows_and_normalizes_business_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "seerfar.xlsx"
            self._workbook(path)
            products = load_seerfar_workbook(path)
        self.assertEqual(len(products), 1)
        self.assertEqual(products[0]["sku"], "1234567890")
        self.assertEqual(products[0]["market_price"], "2536")
        self.assertEqual(products[0]["source_weight"], "950.0")
        self.assertEqual(products[0]["image_url"], "https://ir.ozone.ru/a.jpg")

    def test_saved_mapping_survives_a_fresh_excel_import(self):
        products = [{"id": "123", "market_price": "500", "source_weight": "300"}]
        merged = merge_product_mappings(products, {
            "123": {"supplier_1688_url": "https://detail.1688.com/offer/1.html", "offer_id": "LAMP-1"},
        })
        self.assertEqual(merged[0]["supplier_1688_url"], "https://detail.1688.com/offer/1.html")
        self.assertEqual(merged[0]["offer_id"], "LAMP-1")
        self.assertEqual(merged[0]["price"], "500")
        self.assertEqual(merged[0]["weight"], "300")
        self.assertEqual(mapping_payload(merged)["123"]["offer_id"], "LAMP-1")

    def test_reads_localized_chinese_seerfar_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "seerfar-cn.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append([
                "排名", "主图", "主图", "标题", "详情页地址", "SKU", "品牌", "类目",
                "销售方式", "售价", "销量", "销售额", "毛利率", "评分", "评论数",
                "店铺", "卖家类型", "配送方式", "重量", "上架时间",
            ])
            sheet.append([
                1, '=IF(ISERROR(IMAGE("https://ir.ozone.ru/cn.jpg")),"",IMAGE("https://ir.ozone.ru/cn.jpg"))',
                "https://ir.ozone.ru/cn.jpg", "按摩滚轮", "https://www.ozon.ru/product/3687407905",
                3687407905, "品牌", "运动按摩器", "跨境卖家可售", "212₽", 21514,
                "4390503₽", "49.9%", 4.8, 320, "店铺", "跨境卖家", "FBO", "300 g", "1个月",
            ])
            workbook.save(path)
            products = load_seerfar_workbook(path)

        self.assertEqual(len(products), 1)
        self.assertEqual(products[0]["title"], "按摩滚轮")
        self.assertEqual(products[0]["listing_url"], "https://www.ozon.ru/product/3687407905")
        self.assertEqual(products[0]["sku"], "3687407905")
        self.assertEqual(products[0]["image_url"], "https://ir.ozone.ru/cn.jpg")
        self.assertEqual(products[0]["market_price"], "212")
        self.assertEqual(products[0]["sales"], "21514")
        self.assertEqual(products[0]["source_weight"], "300")

    def test_rejects_unrelated_excel_layout(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.xlsx"
            workbook = Workbook()
            workbook.active.append(["Name", "Value"])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "缺少列"):
                load_seerfar_workbook(path)


if __name__ == "__main__":
    unittest.main()
