"""Historical ledger previews using ERP 1.1, never the old create endpoint."""
import hashlib
import json
import os
import uuid
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from core import atomic_write_json, load_json
from erp_backup import now, no_credentials
from erp_incremental import IncrementalOutbox, text
from product_ledger import LEDGER_COLUMNS, SHEET_TITLE, _LEDGER_LOCK, _process_file_lock

ROW_ID = '_ERP台账记录ID'


def ensure_row_ids(path):
    with _LEDGER_LOCK, _process_file_lock(Path(path).resolve()):
        _ensure_row_ids(path)


def _ensure_row_ids(path):
    path = Path(path)
    workbook = load_workbook(path)
    temporary = path.with_name(path.stem + '.erp-' + uuid.uuid4().hex + '.xlsx')
    try:
        sheet = workbook[SHEET_TITLE]
        headers = [c.value for c in sheet[1]]
        column = headers.index(ROW_ID) + 1 if ROW_ID in headers else sheet.max_column + 1
        sheet.cell(1, column, ROW_ID)
        sheet.column_dimensions[get_column_letter(column)].hidden = True
        seen = set()
        for row in range(2, sheet.max_row + 1):
            if not any(sheet.cell(row, c).value is not None for c in range(1, column)):
                continue
            value = sheet.cell(row, column).value
            if not value:
                value = uuid.uuid4().hex
                sheet.cell(row, column, value)
            if value in seen:
                raise ValueError('台账记录 ID 重复，请核对复制的行')
            seen.add(value)
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def read_ledger(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET_TITLE]
        rows = sheet.iter_rows(values_only=True)
        columns = {c.title: c.key for c in LEDGER_COLUMNS}
        columns[ROW_ID] = '_ledger_identity'
        keys = [columns.get(str(value).strip()) for value in next(rows, ())]
        if not any(keys):
            raise ValueError('未识别产品台账表头')
        for number, values in enumerate(rows, 2):
            record = {key: (value.isoformat() if isinstance(value, (date, datetime)) else value)
                      for key, value in zip(keys, values) if key}
            if any(value is not None and str(value).strip() for key, value in record.items() if not key.startswith('_')):
                yield number, record
    finally:
        workbook.close()


def historical_context(root, row):
    jobs = load_json(Path(root) / 'auto_jobs.json', {}).get('jobs', {})
    if isinstance(jobs, dict):
        jobs = list(jobs.values())
    matches = [job for job in jobs if text(job.get('inputs', {}).get('offer_id')) == text(row.get('offer_id'))
               and text(job.get('inputs', {}).get('ozon_shop_id')) == text(row.get('shop_config_id'))]
    inputs = {}
    if matches and len({job.get('inputs', {}).get('product_group_id') for job in matches}) == 1:
        job = matches[-1]
        inputs = {**job.get('inputs', {}), '_job_id': job['id']}
    submitted = None
    offer = text(row.get('offer_id'))
    if offer and '/' not in offer and '\\' not in offer:
        candidate = load_json(Path(root) / '输出' / ('submitted-' + offer + '.json'), {})
        task_id = row.get('import_task_id') or (matches[-1].get('task_id') if matches else None)
        if candidate.get('item', {}).get('offer_id') == offer and inputs and task_id and text(candidate.get('task_id')) == text(task_id):
            submitted = candidate
    return inputs, submitted


def prepare_ledger(rows, root, target, log=lambda message: None):
    rows = list(rows)
    root = Path(root)
    outbox = IncrementalOutbox(root / 'erp-backup')
    legacy = load_json(root / 'erp-ledger' / 'state.json', {}).get('records', {})
    report = {'started_at': now(), 'items': []}
    for number, row in rows:
        result = {'row': number, 'offer_id': row.get('offer_id')}
        try:
            if text(row.get('platform')).casefold() != 'ozon' or text(row.get('listing_mode')) not in {'跟卖', 'follow'}:
                raise ValueError('只上传明确为 Ozon 跟卖的台账；其他平台/模式或来源未知不冒充 Ozon')
            locator = [target, text(row.get('platform')), text(row.get('shop_config_id') or row.get('shop_name')), text(row.get('offer_id'))]
            old_key = hashlib.sha256(json.dumps(locator, ensure_ascii=False).encode()).hexdigest()
            old = legacy.get(old_key, {})
            if old.get('status') in {'completed', 'sending'}:
                result['legacy_product_id'] = old.get('product_id')
                raise ValueError('旧普通草稿已创建或结果待核对，需 ERP 迁移绑定；不会新建重复产品')
            inputs, submitted = historical_context(root, row)
            path = outbox.prepare_row(row, inputs, target, submitted)
            result.update(status='prepared', manifest_path=str(path))
        except Exception as error:
            result.update(status='failed', error=str(error))
            log(f'台账第 {number} 行：{error}')
        report['items'].append(result)
    buckets = {}
    for number, row in rows:
        for key in ('supplier_1688_url', 'ozon_source_url', 'title'):
            value = text(row.get(key))
            if value:
                buckets.setdefault((key, value), []).append(number)
    suspected = [{'rows': numbers, 'reason': '仅标题或链接相同，不构成同一可售规格依据；不会自动合并'}
                 for numbers in buckets.values() if len(numbers) > 1]
    report['suspected_duplicates'] = suspected
    if suspected:
        log(f'发现 {len(suspected)} 组疑似重复，仅供人工核对，不自动合并；详见上传报告')
    atomic_write_json(root / 'erp-ledger' / 'latest-report.json', report)
    return report


def preview_ledger(rows, root, client, log=lambda text: None):
    report = prepare_ledger(rows, root, client.url, log)
    client.login()
    for item in report['items']:
        if item['status'] != 'prepared':
            continue
        try:
            path = Path(item['manifest_path'])
            preview = client.call('POST', '/api/v1/catalog/follow-backups/validate', data=path.read_bytes(),
                                  headers={'Content-Type': 'application/json; charset=utf-8'})
            no_credentials(preview)
            atomic_write_json(path.with_suffix('.preview.json'), preview)
            item['preview'] = preview
            log('台账预览：' + json.dumps({'row': item['row'], 'items': preview.get('items'), 'warnings': preview.get('warnings')}, ensure_ascii=False))
            item['status'] = 'ready' if preview.get('valid') and (preview.get('can_import') or preview.get('replayed')) else 'failed'
        except Exception as error:
            item.update(status='failed', error=str(error))
    atomic_write_json(Path(root) / 'erp-ledger' / 'latest-report.json', report)
    return report


def upload_ledger_report(report, root, client, log=lambda text: None):
    for item in report['items']:
        if item['status'] != 'ready':
            continue
        try:
            receipt = client.upload(item['manifest_path'], log)
            conflicts = [conflict for entry in receipt.get('items', []) for conflict in entry.get('field_conflicts', [])]
            item.update(status='conflicts' if conflicts else 'completed', receipt=receipt)
            log(f"台账第 {item['row']} 行：" + ('已合并，存在资料差异待确认' if conflicts else '资料已保存/合并'))
        except Exception as error:
            item.update(status='pending', error=str(error))
            log(f"台账第 {item['row']} 行待重试：{error}")
        atomic_write_json(Path(root) / 'erp-ledger' / 'latest-report.json', report)
    return report
